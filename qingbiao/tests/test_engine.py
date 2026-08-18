from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from docx import Document

from qingbiao.engine.economic import compare_to_limit, cross_compare_prices, parse_excel_bid
from qingbiao.engine.metadata import compare_file_properties, extract_file_properties
from qingbiao.engine.report import build_report
from qingbiao.engine.technical import cross_similar, extract_text, split_paragraphs, validate_one


def _xlsx(path: Path, rows: list, creator: str = "张三") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "分部分项"
    ws.append(["项目编码", "项目名称", "计量单位", "工程量", "综合单价", "合价"])
    for row in rows:
        ws.append(row)
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = creator
    wb.properties.company = "同一电脑测试"
    wb.save(path)
    return path


def test_parse_and_limit_and_cross(tmp_path: Path):
    limit = _xlsx(
        tmp_path / "limit.xlsx",
        [
            ["010101", "挖土方", "m3", 100, 50, 5000],
            ["010102", "C30混凝土", "m3", 20, 400, 8000],
        ],
        creator="招标人",
    )
    a = _xlsx(
        tmp_path / "a.xlsx",
        [
            ["010101", "挖土方", "m3", 100, 50, 5000],
            ["010102", "C30混凝土", "m3", 20, 399.9, 7998],
        ],
        creator="甲公司",
    )
    b = _xlsx(
        tmp_path / "b.xlsx",
        [
            ["010101", "挖土方", "m3", 100, 50, 5000],
            ["010102", "C30混凝土", "m3", 20, 400, 8000],
        ],
        creator="乙公司",
    )
    c = _xlsx(
        tmp_path / "c.xlsx",
        [
            ["010101", "挖土方", "m3", 100, 48, 4800],
            ["010102", "C30混凝土", "m3", 20, 410, 8200],
        ],
        creator="丙公司",
    )
    lim = parse_excel_bid(limit, "限价")
    books = [parse_excel_bid(a, "甲"), parse_excel_bid(b, "乙"), parse_excel_bid(c, "丙")]
    assert lim.items and all(x.items for x in books)
    vs = compare_to_limit(lim, books, similar_pct=0.005, abs_tol=0.01)
    cats = {f["category"] for f in vs}
    assert "与最高投标限价单价相同" in cats
    assert "超过最高投标限价" in cats
    cross = cross_compare_prices(books)
    assert any(f["category"] == "多家投标单价相同" for f in cross)


def test_same_author_metadata(tmp_path: Path):
    p1 = _xlsx(tmp_path / "m1.xlsx", [["010101", "挖土方", "m3", 1, 10, 10]], creator="王五")
    p2 = _xlsx(tmp_path / "m2.xlsx", [["010101", "挖土方", "m3", 1, 11, 11]], creator="王五")
    entries = [
        {"bidder": "甲", "filename": "m1.xlsx", "props": extract_file_properties(p1)},
        {"bidder": "乙", "filename": "m2.xlsx", "props": extract_file_properties(p2)},
    ]
    hits = compare_file_properties(entries)
    assert hits
    assert any("同一账号" in h["category"] or "creator=" in h["detail"] for h in hits)


def test_technical_single_and_cross(tmp_path: Path):
    profile = {"floors": "6层", "area": "3000", "structure": "钢结构"}
    text_bad = (
        "本工程采用混凝土剪力墙作为主体。建筑面积：5000。共18层。"
        "依据 GB 50010-2002 和 GB 50017-2003 设计。出现混泥土浇筑。"
        + "专项施工方案应当结合现场起重机械和钢梁钢柱安装。" * 3
    )
    text_ok = (
        "本工程为钢结构厂房，采用钢梁、钢柱、高强度螺栓连接，压型钢板组合楼盖。"
        "执行 GB 55006-2021 钢结构通用规范。建筑面积：3000。共6层。"
        + "钢梁吊装顺序单独描述以免雷同段落。" * 3
    )
    copy = text_ok
    issues = validate_one(text_bad, profile)
    cats = {i["category"] for i in issues}
    assert "引用过期标准" in cats
    assert "错别字" in cats
    assert "与结构类型不匹配" in cats
    p1 = tmp_path / "t1.docx"
    p2 = tmp_path / "t2.docx"
    p3 = tmp_path / "t3.docx"
    for path, body in [(p1, text_ok), (p2, copy), (p3, text_bad)]:
        d = Document()
        d.add_paragraph(body)
        d.save(path)
    docs = []
    for path, name in [(p1, "甲"), (p2, "乙"), (p3, "丙")]:
        t = extract_text(path)
        docs.append({"bidder": name, "text": t, "paragraphs": split_paragraphs(t)})
    cross = cross_similar(docs, threshold=0.86)
    assert any("一致" in f["category"] for f in cross)


def test_report_docx(tmp_path: Path):
    dest = tmp_path / "清标报告.docx"
    path = build_report(
        {"name": "测试项目", "structure": "框架", "floors": "3", "area": "1000"},
        {
            "economic": {
                "bidders": ["甲", "乙", "丙"],
                "limit_file": "limit.xlsx",
                "findings": [
                    {
                        "bidder": "甲",
                        "category": "与最高投标限价单价相同",
                        "severity": "高",
                        "item_code": "010101",
                        "item_name": "挖土方",
                        "detail": "单价 50",
                    }
                ],
            },
            "technical": {
                "bidders": ["甲", "乙", "丙"],
                "single": [{"bidder": "丙", "category": "错别字", "severity": "中", "detail": "混泥土"}],
                "cross": [{"bidder": "甲 / 乙", "category": "全文高度一致", "severity": "高", "detail": "92%"}],
            },
            "metadata": [
                {
                    "bidders": ["甲", "乙"],
                    "category": "疑似同一账号",
                    "severity": "高",
                    "detail": "creator=张三",
                }
            ],
        },
        dest=dest,
    )
    assert path.exists()
    assert path.stat().st_size > 1000
